from __future__ import annotations

from pathlib import Path
from typing import Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "comparativo_herramientas_minimarket.docx"

CONTENT_WIDTH_DXA = 14256

ACCENT = "155E4E"
ACCENT_DARK = "0F3F35"
ACCENT_SOFT = "EAF6F1"
ROSE = "C91546"
ROSE_SOFT = "FFF1F4"
AMBER_SOFT = "FFF6D7"
DARK = "231F20"
MUTED = "6B5B62"
GRID = "D9CDD2"
WHITE = "FFFFFF"
ROW_ALT = "FCFAFB"


def length_to_dxa(length) -> int:
    return int(round(length.twips))


def column_widths_from_weights(weights: Sequence[float], total_width_dxa: int) -> list[int]:
    total_weight = sum(weights)
    widths = [int(round(total_width_dxa * (weight / total_weight))) for weight in weights]
    widths[-1] += total_width_dxa - sum(widths)
    return widths


def ensure_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def set_width(parent, tag: str, width_dxa: int) -> None:
    width = ensure_child(parent, tag)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(width_dxa))


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = ensure_child(tc_pr, "w:tcMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        margin = ensure_child(margins, f"w:{side}")
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def apply_table_geometry(table, column_widths_dxa: Sequence[int], indent_dxa: int = 0) -> None:
    table_width = sum(column_widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    set_width(tbl_pr, "w:tblW", table_width)

    indent = ensure_child(tbl_pr, "w:tblInd")
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), str(indent_dxa))

    layout = ensure_child(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in column_widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for col_idx, width in enumerate(column_widths_dxa):
        table.columns[col_idx].width = Twips(width)

    for row in table.rows:
        for col_idx, cell in enumerate(row.cells):
            cell.width = Twips(column_widths_dxa[col_idx])
            tc_pr = cell._tc.get_or_add_tcPr()
            set_width(tc_pr, "w:tcW", column_widths_dxa[col_idx])
            set_cell_margins(cell)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = ensure_child(tc_pr, "w:shd")
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRID, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = ensure_child(tc_pr, "w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = ensure_child(borders, f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_paragraph_border_bottom(paragraph, color: str = GRID, size: str = "6") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = ensure_child(p_pr, "w:pBdr")
    bottom = ensure_child(borders, "w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def style_run(run, *, bold=False, size=10, color=DARK) -> None:
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def write_cell(cell, text: str, *, bold=False, size=8.5, color=DARK, align=None) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    paragraph.text = ""
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.06
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    style_run(run, bold=bold, size=size, color=color)


def add_styled_paragraph(document: Document, text: str, *, size=10, color=DARK, bold=False):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    style_run(run, size=size, color=color, bold=bold)
    return paragraph


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor.from_string(ACCENT_DARK if level == 1 else ACCENT)


def add_note_table(document: Document, title: str, body: str, fill: str = ACCENT_SOFT) -> None:
    table = document.add_table(rows=1, cols=1)
    apply_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color="E3D4BA")
    set_cell_margins(cell, top=150, bottom=150, start=180, end=180)

    title_paragraph = cell.paragraphs[0]
    title_paragraph.paragraph_format.space_after = Pt(4)
    title_run = title_paragraph.add_run(title)
    style_run(title_run, bold=True, size=10.5, color=ACCENT_DARK)

    body_paragraph = cell.add_paragraph()
    body_paragraph.paragraph_format.space_after = Pt(0)
    body_paragraph.paragraph_format.line_spacing = 1.08
    body_run = body_paragraph.add_run(body)
    style_run(body_run, size=9.2, color=DARK)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    style_run(run, size=8, color=MUTED)


def style_table(table, header_fill: str = ACCENT_DARK) -> None:
    for row_index, row in enumerate(table.rows):
        set_row_cant_split(row)
        if row_index == 0:
            set_repeat_table_header(row)
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_shading(cell, header_fill if row_index == 0 else (WHITE if row_index % 2 else ROW_ALT))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.06
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.1 if row_index else 8.4)
                    run.font.color.rgb = RGBColor.from_string(WHITE if row_index == 0 else DARK)
                    if row_index == 0:
                        run.bold = True


def add_metadata_table(document: Document) -> None:
    values = [
        ("Proyecto", "Minimarket IA"),
        ("Repositorio", "Next.js + Supabase"),
        ("Fecha", "14 de mayo de 2026"),
        ("Documento", "Comparativo de herramientas"),
    ]
    table = document.add_table(rows=2, cols=4)
    apply_table_geometry(table, column_widths_from_weights([1, 1, 1, 1], CONTENT_WIDTH_DXA))
    for index, (label, value) in enumerate(values):
        label_cell = table.cell(0, index)
        value_cell = table.cell(1, index)
        write_cell(label_cell, label, bold=True, size=8.3, color=ROSE, align=WD_ALIGN_PARAGRAPH.CENTER)
        write_cell(value_cell, value, bold=True, size=8.5, color=DARK, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(label_cell, ROSE_SOFT)
        set_cell_shading(value_cell, WHITE)
        set_cell_border(label_cell)
        set_cell_border(value_cell)
    document.add_paragraph()


def add_tools_table(document: Document) -> None:
    headers = [
        "Etapa",
        "Herramienta exacta",
        "Uso en el proyecto",
        "Alternativas equivalentes",
        "Por que conviene mantenerla",
    ]
    rows = [
        [
            "Edicion",
            "Visual Studio Code 1.119.0",
            "IDE principal para editar componentes, rutas API, scripts, SQL, README y documentacion del proyecto.",
            "WebStorm, Visual Studio, Sublime Text, Notepad++",
            "Ya esta instalado, es ligero, tiene soporte fuerte para TypeScript, React, Next.js, Git y extensiones utiles sin cambiar el flujo de trabajo.",
        ],
        [
            "Terminal",
            "Windows PowerShell",
            "Ejecutar comandos de desarrollo, instalacion, build, lint, typecheck, Git y scripts de generacion.",
            "Git Bash, CMD, WSL, terminal de Linux/macOS",
            "Funciona en el entorno actual de Windows y permite operar Node, npm y Git sin agregar una capa extra para este prototipo.",
        ],
        [
            "Versiones",
            "Git 2.54.0.windows.1",
            "Registrar cambios, comparar avances, crear ramas y preparar el codigo para GitHub y despliegues.",
            "Mercurial, SVN, copias manuales por carpetas",
            "Es el estandar del ecosistema web, se integra con VS Code, GitHub y Vercel, y evita perder cambios importantes.",
        ],
        [
            "Repositorio",
            "GitHub",
            "Guardar el repositorio remoto, revisar historial y conectar despliegues automaticos hacia Vercel.",
            "GitLab, Bitbucket, Azure DevOps",
            "Tiene integracion directa con Vercel y es suficiente para control de versiones, revision y entrega del proyecto.",
        ],
        [
            "Runtime JS",
            "Node.js v24.15.0",
            "Ejecutar Next.js, scripts de desarrollo y herramientas npm en local.",
            "Bun, Deno, otros runtimes Node LTS",
            "Es el runtime natural del stack usado. Mantiene compatibilidad con Next.js, npm y las dependencias declaradas.",
        ],
        [
            "Paquetes",
            "npm v11.12.1 + package-lock.json",
            "Instalar dependencias y ejecutar comandos: dev, build, lint y typecheck.",
            "pnpm, Yarn, Bun install",
            "El proyecto ya esta bloqueado con package-lock.json, asi que npm reduce cambios innecesarios y mantiene instalaciones reproducibles.",
        ],
        [
            "Framework web",
            "Next.js 16.2.4 con App Router",
            "Construir la web, el panel administrativo y las rutas API internas y moviles desde un solo repositorio.",
            "Vite + Express, Laravel, Django, Astro, Angular + NestJS",
            "Une frontend y backend, encaja con Vercel, reduce configuracion y ya coincide con la arquitectura existente del proyecto.",
        ],
        [
            "Interfaz UI",
            "React 19.2.4",
            "Crear componentes reutilizables como dashboard, chatbot, formularios, catalogo y vistas administrativas.",
            "Vue, Angular, Svelte, HTML/JavaScript manual",
            "El proyecto ya usa componentes React. Cambiar de libreria implicaria rehacer pantallas sin una mejora proporcional.",
        ],
        [
            "Lenguaje",
            "TypeScript 5",
            "Tipar productos, categorias, facturas, respuestas de API y estructuras compartidas del sistema.",
            "JavaScript puro, PropTypes, validaciones manuales",
            "Detecta errores antes de ejecutar, mejora el mantenimiento y protege contratos entre componentes, librerias y rutas API.",
        ],
        [
            "Estilos",
            "Tailwind CSS 4.2.4 + @tailwindcss/postcss 4.2.4",
            "Diseñar pantallas responsivas, formularios, tablas, botones y estados visuales desde clases utilitarias.",
            "Bootstrap, Material UI, Sass, CSS Modules",
            "Permite avanzar rapido sin crear un sistema CSS grande. Ademas conserva control visual propio del minimarket.",
        ],
        [
            "Iconos",
            "lucide-react 1.14.0",
            "Mostrar iconos consistentes para acciones como buscar, subir, descargar, editar, eliminar y navegar.",
            "Font Awesome, Material Icons, SVG manual",
            "Es liviano, moderno y se integra directo con React. Evita mezclar estilos visuales o dibujar iconos uno por uno.",
        ],
        [
            "Calidad",
            "ESLint 9 + eslint-config-next 16.2.4 + tsc --noEmit",
            "Revisar reglas del proyecto, errores de tipado y problemas antes de compilar o desplegar.",
            "Biome, Prettier solamente, revision manual",
            "Usa reglas oficiales de Next.js y se combina con TypeScript. La revision manual por si sola no detecta suficientes fallos.",
        ],
        [
            "Backend/API",
            "Route Handlers de Next.js en app/api",
            "Atender login admin, catalogo, productos, categorias, facturas, uploads, chatbot y endpoints moviles.",
            "Express separado, NestJS, Supabase Edge Functions, servidor Laravel",
            "Mantiene las APIs junto a la app web, simplifica despliegue y evita administrar otro servidor para el prototipo.",
        ],
        [
            "Base de datos",
            "Supabase con PostgreSQL",
            "Guardar catalogo, categorias, productos, pedidos, quejas, facturas y datos relacionales.",
            "Firebase, MySQL en VPS, MongoDB, SQLite, archivos JSON",
            "PostgreSQL encaja mejor con datos relacionales. Supabase agrega panel, API, storage y seguridad sin montar infraestructura propia.",
        ],
        [
            "SDK Supabase",
            "@supabase/supabase-js 2.105.1",
            "Consultar, insertar, actualizar y eliminar informacion de Supabase desde el servidor y utilidades del proyecto.",
            "Fetch REST directo, Prisma, Drizzle ORM",
            "Es el SDK oficial, reduce codigo repetido y mantiene acceso simple a base de datos, storage y servicios Supabase.",
        ],
        [
            "Modelo SQL",
            "supabase/schema.sql + supabase/seed.sql",
            "Definir tablas, relaciones, datos base y reinicio del entorno de prueba.",
            "Migraciones ORM, scripts manuales en consola, hojas Excel como base unica",
            "SQL es directo para PostgreSQL y deja el modelo versionado dentro del repositorio junto al codigo.",
        ],
        [
            "Imagenes",
            "Supabase Storage, bucket product-images",
            "Guardar imagenes de productos subidas desde el panel administrativo.",
            "Cloudinary, Amazon S3, Firebase Storage, carpeta public local",
            "Se integra con Supabase y evita un servicio adicional. La carpeta local no es buena para produccion o multiusuario.",
        ],
        [
            "Auth admin",
            "Ruta /api/admin/login + variables de entorno; Supabase Auth como evolucion posible",
            "Validar acceso del administrador sin exponer llaves secretas en el navegador.",
            "Auth0, Firebase Auth, NextAuth, usuarios manuales en frontend",
            "Para prototipo es simple y controlado. Las llaves sensibles viven en servidor o Vercel, no en variables NEXT_PUBLIC.",
        ],
        [
            "Excel",
            "xlsx-js-style 1.2.0",
            "Leer plantillas, importar productos, exportar catalogos con hojas, estilos, filtros, resumen y anchos de columna.",
            "ExcelJS, SheetJS xlsx, CSV/PapaParse, Google Sheets API",
            "CSV pierde hojas y formato. Google Sheets agrega permisos externos. xlsx-js-style permite generar archivos Excel completos y presentables.",
        ],
        [
            "Documentacion",
            "README.md + docs/*.md + DOCX generado con Python/python-docx",
            "Explicar instalacion, variables, rutas, APIs Android y decisiones tecnicas del proyecto.",
            "Notion, Confluence, Google Docs, documentos Word manuales",
            "La documentacion versionada vive con el codigo. El DOCX sirve para entregar evidencia formal sin separar la informacion del proyecto.",
        ],
        [
            "Pruebas",
            "npm run dev, navegador, npm run typecheck, npm run lint, npm run build",
            "Probar la app en localhost, validar tipos, revisar reglas y confirmar compilacion de produccion.",
            "Pruebas manuales solamente, Cypress, Playwright, Jest/Vitest",
            "Para el alcance actual cubre la barrera minima. Cuando el sistema crezca, se puede sumar Playwright o unit tests sin reemplazar el stack.",
        ],
        [
            "Despliegue",
            "Vercel",
            "Publicar la app Next.js, configurar variables de entorno y conectar con GitHub para despliegues.",
            "Netlify, Render, Railway, VPS tradicional",
            "Vercel es la opcion mas directa para Next.js. Reduce configuracion de servidor y maneja builds, dominios y variables facilmente.",
        ],
        [
            "Secretos",
            ".env.local, .env.example y variables de entorno en Vercel",
            "Separar URLs, llaves anon, service role y credenciales administrativas del codigo fuente.",
            "Valores hardcodeados, archivos JSON privados, configuracion manual en codigo",
            "Evita filtrar secretos y permite cambiar entre local, demo y produccion sin tocar la logica de la app.",
        ],
        [
            "Movil futuro",
            "APIs /api/mobile/*; Android Studio si se desarrolla app nativa",
            "Permitir que una app Android consuma catalogo, pedidos, quejas y devoluciones desde la URL desplegada.",
            "Flutter, React Native, Ionic, PWA web",
            "El repo actual ya deja APIs listas. Android Studio conviene si se busca app nativa; Flutter o React Native son alternativas si se desea multiplataforma.",
        ],
    ]

    table = document.add_table(rows=1, cols=len(headers))
    column_widths = column_widths_from_weights([1.1, 1.55, 2.05, 1.7, 3.5], CONTENT_WIDTH_DXA)
    apply_table_geometry(table, column_widths)
    for index, header in enumerate(headers):
        write_cell(table.rows[0].cells[index], header, bold=True, size=8.3, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            write_cell(cells[index], value, bold=index in (0, 1), size=7.7)

    apply_table_geometry(table, column_widths)
    style_table(table)


def add_decision_summary(document: Document) -> None:
    add_heading(document, "Decision recomendada", 1)
    conclusions = [
        "Mantener el stack actual es lo mas conveniente porque ya cubre desarrollo, base de datos, APIs, Excel, documentacion y despliegue sin dividir el proyecto en demasiadas tecnologias.",
        "Las alternativas son validas, pero en este caso aumentarian configuracion, aprendizaje o mantenimiento sin resolver mejor el objetivo principal del minimarket.",
        "Para crecimiento futuro, las mejoras naturales serian agregar pruebas automatizadas con Playwright/Vitest y formalizar la app Android en un repositorio propio.",
    ]
    for item in conclusions:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(item)
        style_run(run, size=9.5, color=DARK)

    add_note_table(
        document,
        "Resumen final",
        "La mejor ruta es seguir con Visual Studio Code, Git/GitHub, Node.js, npm, Next.js, React, TypeScript, Tailwind CSS, Supabase, xlsx-js-style y Vercel. Esta combinacion es suficiente para construir, probar, documentar, desplegar y ampliar el sistema de Minimarket IA de principio a fin.",
        fill=AMBER_SOFT,
    )


def build_document() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    actual_width = length_to_dxa(section.page_width) - length_to_dxa(section.left_margin) - length_to_dxa(section.right_margin)
    if actual_width != CONTENT_WIDTH_DXA:
        raise RuntimeError(f"Unexpected content width: {actual_width}")

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(ACCENT_DARK if style_name != "Title" else DARK)

    header = section.header.paragraphs[0]
    header.text = ""
    header_run = header.add_run("Minimarket IA | Comparativo de herramientas")
    style_run(header_run, size=8.5, color=MUTED)
    set_paragraph_border_bottom(header)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("Pagina ")
    style_run(footer_run, size=8, color=MUTED)
    add_page_number(footer)

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(2)
    title.add_run("Tabla comparativa de herramientas del proyecto")

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    subtitle_run = subtitle.add_run(
        "Herramientas exactas que se usaran en Minimarket IA, alternativas posibles y justificacion de la eleccion."
    )
    style_run(subtitle_run, size=10.5, color=MUTED)

    add_metadata_table(document)

    add_note_table(
        document,
        "Criterio de seleccion",
        "La comparacion prioriza herramientas ya presentes en el repositorio o indicadas en el README: menor friccion, compatibilidad con Next.js/Supabase, facilidad de despliegue, mantenimiento para un prototipo de titulacion y posibilidad de crecer hacia una app Android.",
        fill=ACCENT_SOFT,
    )

    add_heading(document, "Tabla principal", 1)
    add_tools_table(document)
    document.add_paragraph()
    add_decision_summary(document)

    document.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
